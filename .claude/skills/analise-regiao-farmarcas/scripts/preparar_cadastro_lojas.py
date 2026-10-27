"""Converte o cadastro de lojas (.xlsx) para CSV, uma vez.

Uso (rodar a partir da raiz do projeto, SEM -I, porque precisa de openpyxl):
    python .claude/skills/analise-regiao-farmarcas/scripts/preparar_cadastro_lojas.py

Por que existe: `openpyxl` está no site-packages de usuário e não é visível sob
`python -I`, que é como os scripts de análise rodam. Em vez de afrouxar o
isolamento de todo mundo, convertemos uma vez e os scripts de análise leem CSV
com a biblioteca padrão.

Saída: `cache/cadastro_lojas.csv` dentro da pasta da skill (fora do git).
Rodar de novo sempre que o cadastro for reexportado.
"""

from __future__ import annotations

import csv
import os
import sys

COLUNAS = [
    "Id",
    "Data de Inauguração",
    "Rede",
    "Razão Social",
    "Nome Fantasia",
    "CNPJ",
    "Bandeira",
    "Modelo de Negócio",
    "Tipo",
    "Status",
    "Situação",
    "Data de Desativação",
    "Cidade - Endereço Receita",
    "Estado - Endereço Receita",
    "Bairro - Endereço Receita",
    "Latitude - Endereço Receita",
    "Longitude - Endereço Receita",
    "Nome do responsável pelo GE",
]

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, "cache", "cadastro_lojas.csv")


def achar_xlsx(base_dir: str) -> str:
    import glob

    hits = sorted(glob.glob(os.path.join(base_dir, "relatorio-lojas-*.xlsx")))
    if not hits:
        raise FileNotFoundError(f"nenhum relatorio-lojas-*.xlsx em {base_dir!r}")
    return hits[-1]


def main() -> int:
    try:
        import openpyxl
    except ImportError:
        print(
            "openpyxl não encontrado. Rode SEM a flag -I, ou instale com:\n"
            "    python -m pip install --user openpyxl",
            file=sys.stderr,
        )
        return 1

    base_dir = os.environ.get("ORBITA_BASES", r"C:\Users\zequi\Downloads\bases órbita")
    caminho = sys.argv[1] if len(sys.argv) > 1 else achar_xlsx(base_dir)

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = next(linhas)
    indice = {nome: i for i, nome in enumerate(cabecalho)}

    faltando = [c for c in COLUNAS if c not in indice]
    if faltando:
        print(f"colunas ausentes no xlsx: {faltando}", file=sys.stderr)
        return 1

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    total = 0
    with open(SAIDA, "w", encoding="utf-8", newline="") as fh:
        escritor = csv.writer(fh)
        escritor.writerow(COLUNAS)
        for linha in linhas:
            if linha is None or all(v is None for v in linha):
                continue
            escritor.writerow(
                ["" if linha[indice[c]] is None else str(linha[indice[c]]).strip() for c in COLUNAS]
            )
            total += 1

    print(f"origem : {caminho}")
    print(f"destino: {SAIDA}")
    print(f"lojas  : {total}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
