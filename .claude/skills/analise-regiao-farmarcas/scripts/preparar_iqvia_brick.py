"""Converte a planilha de faturamento da IQVIA por brick para CSV, uma vez.

Uso (a partir da raiz do projeto, SEM -I, porque precisa de openpyxl):
    python .claude/skills/analise-regiao-farmarcas/scripts/preparar_iqvia_brick.py

Entrada: `Faturamento IQVIA - <MAT> - Detalhado por CNPJ, BRICK e UTC*.xlsx`,
aba `DADOS POR BRICK E UTC`. É a única fonte do dump que traz **unidades
vendidas** (`unidade`) junto com a **receita a preço de consumidor**
(`preco_consumidor`) — e é por isso que ela destrava o preço médio.

    preço médio por unidade = preco_consumidor / unidade

Atenção: `preco_consumidor` é RECEITA, não preço unitário. O nome engana.

Saída: `cache/iqvia_brick.csv`. Rodar de novo quando a planilha for atualizada.
"""

from __future__ import annotations

import csv
import glob
import os
import sys

COLUNAS = [
    "flag_mat",
    "anomes",
    "descricao_brick_rpe",
    "codigo_uf",
    "provedor_descricao_grupo",
    "codigo_cnpj",
    "unidade",
    "preco_consumidor",
]

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, "cache", "iqvia_brick.csv")
ABA = "DADOS POR BRICK E UTC"


def main() -> int:
    try:
        import openpyxl
    except ImportError:
        print("openpyxl não encontrado. Rode SEM -I, ou: python -m pip install --user openpyxl", file=sys.stderr)
        return 1

    base_dir = os.environ.get("ORBITA_BASES", r"C:\Users\zequi\Downloads\bases órbita")
    if len(sys.argv) > 1:
        caminho = sys.argv[1]
    else:
        hits = sorted(glob.glob(os.path.join(base_dir, "Faturamento IQVIA*BRICK*.xlsx")))
        if not hits:
            print(f"nenhuma planilha 'Faturamento IQVIA*BRICK*.xlsx' em {base_dir!r}", file=sys.stderr)
            return 1
        # As cópias "(3)" e "(4)" são idênticas; qualquer uma serve.
        caminho = hits[0]

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    if ABA not in wb.sheetnames:
        print(f"aba {ABA!r} não encontrada. Abas: {wb.sheetnames}", file=sys.stderr)
        return 1
    ws = wb[ABA]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = next(linhas)
    indice = {str(nome): i for i, nome in enumerate(cabecalho)}

    faltando = [c for c in COLUNAS if c not in indice]
    if faltando:
        print(f"colunas ausentes: {faltando}", file=sys.stderr)
        return 1

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    total = 0
    with open(SAIDA, "w", encoding="utf-8", newline="") as fh:
        escritor = csv.writer(fh)
        escritor.writerow(COLUNAS)
        for linha in linhas:
            if linha is None or all(v is None for v in linha):
                continue
            escritor.writerow(
                ["" if linha[indice[c]] is None else str(linha[indice[c]]).strip() for c in COLUNAS]
            )
            total += 1

    print(f"origem : {os.path.basename(caminho)}")
    print(f"destino: {SAIDA}")
    print(f"linhas : {total}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
