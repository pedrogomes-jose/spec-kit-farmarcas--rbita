"""Converte as abas de canal da planilha da IQVIA para CSV, uma vez.

Uso (da raiz do projeto, SEM -I, porque precisa de openpyxl):
    python .claude/skills/analise-regiao-farmarcas/scripts/preparar_iqvia_canal.py

Por que importa: é a única fonte **medida** de faturamento de concorrente. O
campo de faturamento da camada Concorrentes é estimado e está preenchido em
menos da metade dos registros; estas duas abas trazem, por **canal e UF**:

    DADOS POR CANAL - CPP  -> receita (soma_cpp) e unidades, visão MAT
    DADOS POR CANAL - PDV  -> quantidade de PDVs, visão YTD

    faturamento médio por loja = receita MAT do canal ÷ PDVs do canal ÷ 12

Os canais são ABRAFARMA, FEBRAFAR, ASSOCIACOES + FRANQUIAS, INDEPENDENTES,
OUTRAS REDES CORPORATIVAS e GRUPO FARMARCAS — e a camada Concorrentes
classifica cada loja numa categoria que mapeia para eles.

Saídas: `cache/iqvia_canal_cpp.csv` e `cache/iqvia_canal_pdv.csv`.
"""

from __future__ import annotations

import csv
import glob
import os
import sys

AQUI = os.path.dirname(os.path.abspath(__file__))
ABAS = {
    "DADOS POR CANAL - CPP": (
        "iqvia_canal_cpp.csv",
        ["visao", "label", "grupo", "bandeira", "setor_nec_aberto", "uf", "soma_cpp", "soma_unidade"],
    ),
    "DADOS POR CANAL - PDV": (
        "iqvia_canal_pdv.csv",
        ["visao", "label", "grupo", "bandeira", "uf", "qtd_pdvs"],
    ),
}


def main() -> int:
    try:
        import openpyxl
    except ImportError:
        print("openpyxl não encontrado. Rode SEM -I, ou: python -m pip install --user openpyxl", file=sys.stderr)
        return 1

    base_dir = os.environ.get("ORBITA_BASES", r"C:\Users\zequi\Downloads\bases órbita")
    caminho = sys.argv[1] if len(sys.argv) > 1 else None
    if not caminho:
        hits = sorted(glob.glob(os.path.join(base_dir, "Faturamento IQVIA*BRICK*.xlsx")))
        if not hits:
            print(f"nenhuma planilha 'Faturamento IQVIA*BRICK*.xlsx' em {base_dir!r}", file=sys.stderr)
            return 1
        caminho = hits[0]

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    os.makedirs(os.path.join(AQUI, "cache"), exist_ok=True)

    for aba, (saida, colunas) in ABAS.items():
        if aba not in wb.sheetnames:
            print(f"aba {aba!r} não encontrada. Abas: {wb.sheetnames}", file=sys.stderr)
            return 1
        ws = wb[aba]
        linhas = ws.iter_rows(values_only=True)
        cabecalho = [str(c) for c in next(linhas)]
        indice = {nome: i for i, nome in enumerate(cabecalho)}
        faltando = [c for c in colunas if c not in indice]
        if faltando:
            print(f"colunas ausentes em {aba}: {faltando}", file=sys.stderr)
            return 1

        destino = os.path.join(AQUI, "cache", saida)
        total = 0
        with open(destino, "w", encoding="utf-8", newline="") as fh:
            escritor = csv.writer(fh)
            escritor.writerow(colunas)
            for linha in linhas:
                if linha is None or all(v is None for v in linha):
                    continue
                escritor.writerow(
                    ["" if linha[indice[c]] is None else str(linha[indice[c]]).strip() for c in colunas]
                )
                total += 1
        print(f"{aba}: {total} linhas -> {destino}")

    wb.close()
    print(f"origem: {os.path.basename(caminho)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
